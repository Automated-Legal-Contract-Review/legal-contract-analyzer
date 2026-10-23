import docx
import pytest

from app import analyzer, chunking, extract, llm, vectors


# ---------------------------------------------------------------- chunking & quote matching
def make_doc(n_pages=10, words=900):
    pages = [f"Page {i} heading.\n\n" + ("Alpha beta gamma delta. " * (words // 4)) for i in range(1, n_pages + 1)]
    return chunking.build_document(pages), pages


def test_chunks_cover_whole_document_with_overlap():
    doc, _ = make_doc()
    chunks = chunking.make_chunks(doc, 3000, 300)
    assert len(chunks) > 3
    assert chunks[0]["start"] == 0
    assert chunks[-1]["end"] == len(doc.full_text)
    for a, b in zip(chunks, chunks[1:]):
        assert b["start"] < a["end"], "chunks must overlap"
        assert b["start"] > a["start"], "chunks must make progress"
    assert all(len(c["text"]) <= 3000 for c in chunks)


def test_chunks_do_not_start_mid_word():
    doc, _ = make_doc()
    for chunk in chunking.make_chunks(doc, 3000, 300)[1:]:
        assert chunk["text"][0] != " "
        previous = doc.full_text[chunk["start"] - 1]
        assert previous in " \n", "a chunk should begin right after whitespace"


def test_page_mapping():
    doc, pages = make_doc(5)
    assert doc.page_count == 5
    assert doc.page_at(0) == 1
    assert doc.page_at(len(doc.full_text) - 1) == 5
    marker = doc.full_text.index("Page 3 heading")
    assert doc.page_at(marker) == 3


def test_matcher_tolerates_formatting_differences():
    text = "The Supplier’s liability\nshall be   UNLIMITED — without any cap whatsoever."
    matcher = chunking.Matcher("Intro text. " + text + " More text.")
    start, end = matcher.find_loose("the supplier's liability shall be unlimited - without any cap")
    assert "Supplier" in matcher.text[start:start + 20]
    assert matcher.find_loose("This sentence is not in the document at all.") is None
    assert matcher.find_loose("too short") is None


def test_matcher_accepts_ellipsis_quotes():
    matcher = chunking.Matcher("The customer shall pay all fees within ten days and this payment is non refundable in every case.")
    assert matcher.find_loose("The customer shall pay all fees ... payment is non refundable in every case") is not None


# ---------------------------------------------------------------- JSON recovery
def test_parse_plain_and_fenced_json():
    assert llm.parse_json_lenient('{"a": 1}') == {"a": 1}
    assert llm.parse_json_lenient('```json\n{"a": 2}\n```') == {"a": 2}
    assert llm.parse_json_lenient('Here you go: {"a": 3} thanks') == {"a": 3}


def test_salvage_truncated_findings():
    truncated = (
        '{"chunk_summary": "Covers payment.", "findings": [{"title": "A", "evidence": "x"}, '
        '{"title": "B", "evidence": "y"}, {"title": "C", "evid'
    )
    result = llm.parse_json_lenient(truncated)
    assert result["chunk_summary"] == "Covers payment."
    assert [f["title"] for f in result["findings"]] == ["A", "B"]


def test_unusable_output_raises():
    with pytest.raises(llm.LLMError):
        llm.parse_json_lenient("I cannot help with that.")


# ---------------------------------------------------------------- findings, dedupe, risk
def finding(level="High", severity=4, start=0, end=50, title="t", category="Liability"):
    return {"risk_level": level, "severity_score": severity, "char_start": start, "char_end": end,
            "title": title, "category": category, "explanation": "x" * 10, "chunk_index": 1}


def test_risk_score_properties():
    assert analyzer.risk_score([]) == 0
    one = analyzer.risk_score([finding()], pages=5)
    many = analyzer.risk_score([finding()] * 6, pages=5)
    assert 0 < one < many <= 100
    assert analyzer.risk_score([finding("Info", 1)] * 50, pages=5) == 0
    assert analyzer.risk_score([finding()] * 5, pages=300) < analyzer.risk_score([finding()] * 5, pages=5)
    assert analyzer.risk_label(0) == "Low" and analyzer.risk_label(90) == "High"


def test_dedupe_merges_overlap_duplicates_and_keeps_distinct():
    a = finding(start=100, end=200, title="Unlimited liability", severity=4)
    b = finding(start=120, end=210, title="Liability not capped", severity=5)  # same clause, other chunk
    c = finding(start=5000, end=5100, title="Unlimited liability", severity=4)  # same title, far away
    result = analyzer.dedupe([a, b, c])
    assert len(result) == 2
    assert max(f["severity_score"] for f in result if f["char_start"] < 1000) == 5


def test_grounded_findings_drops_invented_quotes_and_fixes_severity():
    doc = chunking.build_document(["The supplier accepts unlimited liability for all losses under this agreement."])
    chunk = chunking.make_chunks(doc, 5000, 0)[0]
    raw = [
        {"title": "Real", "category": "liability", "risk_level": "High", "severity_score": 1,
         "evidence": "supplier accepts unlimited liability for all losses", "explanation": "Big exposure", "recommendation": "Cap it"},
        {"title": "Fake", "category": "Other", "risk_level": "High", "severity_score": 5,
         "evidence": "the supplier will gift the customer a castle", "explanation": "Invented", "recommendation": ""},
        "not a dict",
    ]
    result = analyzer.grounded_findings(raw, chunk, doc)
    assert len(result) == 1
    assert result[0]["category"] == "Liability"
    assert result[0]["severity_score"] == 4  # High must be 4-5
    assert result[0]["page_start"] == 1


# ---------------------------------------------------------------- keyword retrieval
def test_keyword_search_ranks_relevant_passage_first():
    passages = [
        {"i": 1, "text": "Payment is due within thirty days of the invoice date.", "page_start": 1, "page_end": 1},
        {"i": 2, "text": "The agreement renews automatically every year unless notice is given.", "page_start": 2, "page_end": 2},
        {"i": 3, "text": "Governing law is the law of England.", "page_start": 3, "page_end": 3},
    ]
    hits = vectors.search_keywords(passages, "When does the agreement renew automatically?")
    assert hits and hits[0]["page_start"] == 2
    assert vectors.search_keywords(passages, "zebra quantum") == []


# ---------------------------------------------------------------- extraction
def test_docx_extraction_includes_tables_in_order(tmp_path):
    path = tmp_path / "c.docx"
    document = docx.Document()
    document.add_paragraph("Services Agreement preamble paragraph with enough words to count as text.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Fee"
    table.cell(0, 1).text = "Ten thousand dollars per month"
    table.cell(1, 0).text = "Term"
    table.cell(1, 1).text = "Twelve months"
    document.add_paragraph("Closing paragraph about termination for convenience on thirty days notice.")
    document.save(path)
    result = extract.extract(path)
    text = "\n".join(result.pages)
    assert "Fee | Ten thousand dollars per month" in text
    assert text.index("preamble") < text.index("Fee |") < text.index("Closing")
    assert result.page_kind == "approx"


def test_txt_encodings(tmp_path):
    body = "Zahlungsbedingungen: Die Zahlung erfolgt innerhalb von 30 Tagen nach Rechnungsdatum. Preis in € netto."
    for name, raw in {
        "utf8.txt": body.encode("utf-8"),
        "cp1252.txt": body.encode("cp1252"),
        "utf16.txt": b"\xff\xfe" + body.encode("utf-16-le"),
    }.items():
        path = tmp_path / name
        path.write_bytes(raw)
        assert "Zahlungsbedingungen" in extract.extract(path).pages[0], name
        assert "€" in extract.extract(path).pages[0], name


def test_bad_files_give_clear_errors(tmp_path):
    fake_pdf = tmp_path / "x.pdf"
    fake_pdf.write_bytes(b"this is not a pdf at all" * 20)
    with pytest.raises(extract.ExtractionError, match="not a valid PDF"):
        extract.extract(fake_pdf)
    fake_docx = tmp_path / "x.docx"
    fake_docx.write_bytes(b"plain text pretending to be docx")
    with pytest.raises(extract.ExtractionError, match="not a valid Word"):
        extract.extract(fake_docx)
    empty = tmp_path / "e.txt"
    empty.write_bytes(b"")
    with pytest.raises(extract.ExtractionError, match="empty"):
        extract.extract(empty)
    tiny = tmp_path / "t.txt"
    tiny.write_text("too short")
    with pytest.raises(extract.ExtractionError, match="Too little"):
        extract.extract(tiny)


# ---------------------------------------------------------------- scanned PDFs
def _scanned_pdf(path):
    import pymupdf
    src = pymupdf.open()
    page = src.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 700), "The Supplier shall have unlimited liability for all losses. Payment is due within thirty days of the invoice date.", fontsize=14)
    pix = page.get_pixmap(dpi=200)
    scan = pymupdf.open()
    scan.new_page().insert_image(scan[0].rect, pixmap=pix)
    scan.save(path)


def test_scanned_pdf_without_ocr_gives_clear_message(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    _scanned_pdf(path)
    monkeypatch.setattr(extract, "ocr_available", lambda: False)
    with pytest.raises(extract.ExtractionError, match="scanned"):
        extract.extract(path)


@pytest.mark.skipif(not extract.ocr_available(), reason="Tesseract not installed")
def test_scanned_pdf_is_read_with_ocr(tmp_path):
    path = tmp_path / "scan.pdf"
    _scanned_pdf(path)
    result = extract.extract(path)
    assert "liability" in result.pages[0].lower()
    assert any("OCR" in note for note in result.notes)


def test_cancelled_queued_analysis_is_not_restarted(monkeypatch):
    """A stale queue entry must not reset a cancelled contract to processing."""
    cancelled_record = {
        "id": "cancelled-id",
        "status": "cancelled",
        "model": "llama3.2",
        "perspective": "neutral",
        "file_path": "does-not-matter.pdf",
    }
    monkeypatch.setattr(analyzer.db, "get", lambda _contract_id: cancelled_record)

    def unexpected(*_args, **_kwargs):
        raise AssertionError("cancelled job should not start model work")

    monkeypatch.setattr(analyzer.llm, "check_ready", unexpected)
    analyzer.run_analysis("cancelled-id")


def test_completed_analysis_is_not_run_twice_from_stale_queue(monkeypatch):
    completed_record = {"id": "done-id", "status": "completed"}
    monkeypatch.setattr(analyzer.db, "get", lambda _contract_id: completed_record)

    def unexpected(*_args, **_kwargs):
        raise AssertionError("completed job should not start model work")

    monkeypatch.setattr(analyzer.llm, "check_ready", unexpected)
    analyzer.run_analysis("done-id")
